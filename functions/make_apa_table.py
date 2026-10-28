# -*- coding: utf-8 -*-
"""
Create APA-7 style patient-features table from TONI preprocessed data.

Requirements:
  pip install pandas numpy python-docx

Run:
  python make_patient_table_apa7.py
"""

from __future__ import annotations

import sys
from pathlib import Path
import re
import ast

import numpy as np
import pandas as pd

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ---- make project root importable (robust) ----
PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from server_config import DATA_PATH  # noqa: E402


# =============================================================================
# ICD parsing + Table-2 mapping (multi-label primary dx, one-row-per-patient)
# =============================================================================
ICD_RE = re.compile(r"[A-Z][0-9]{2}(?:\.[0-9]{1,3})?")

TABLE2_CATS = {
    "substance":  "Mental and behavioral disorders due to substance use",
    "psychotic":  "Schizophrenia, schizotypal and delusional disorders",
    "affective":  "Affective disorders",
    "anxiety":    "Anxiety disorders",
    "ocd":        "Obsessive–compulsive disorder",
    "asr_ptsd":   "Acute stress reaction and posttraumatic stress disorder",
    "adjustment": "Adjustment disorder",
    "somatoform": "Somatoform disorders",
    "eating":     "Eating disorders",
    "personality":"Personality disorders",
    "adhd":       "Hyperkinetic and attention-deficit disorders",
    "other":      "Other disorders",
    "missing":    "Missing diagnosis",
}

TABLE2_ORDER = [
    "Mental and behavioral disorders due to substance use",
    "Schizophrenia, schizotypal and delusional disorders",
    "Affective disorders",
    "Anxiety disorders",
    "Obsessive–compulsive disorder",
    "Acute stress reaction and posttraumatic stress disorder",
    "Adjustment disorder",
    "Somatoform disorders",
    "Eating disorders",
    "Personality disorders",
    "Hyperkinetic and attention-deficit disorders",
    "Other disorders",
    "Missing diagnosis",
]


def parse_icd_codes(x) -> list[str]:
    """Return list of ICD-like codes from x, robust to malformed strings."""
    if x is None or x is pd.NA or (isinstance(x, float) and pd.isna(x)):
        return []

    if isinstance(x, list):
        out = []
        for c in x:
            s = str(c).strip().upper()
            if s and s not in {"NONE", "NAN", "NULL", "<NA>"}:
                out.append(s)
        return out

    if isinstance(x, str):
        s = x.strip()
        if s in ("", "[]", ","):
            return []
        if s.upper() in {"NONE", "NAN", "NULL", "<NA>"}:
            return []
        try:
            obj = ast.literal_eval(s)
            if isinstance(obj, list):
                out = []
                for c in obj:
                    cc = str(c).strip().upper()
                    if cc and cc not in {"NONE", "NAN", "NULL", "<NA>"}:
                        out.append(cc)
                return out
        except (ValueError, SyntaxError):
            pass
        return ICD_RE.findall(s.upper())

    return []


def norm_icd(code, level=3) -> str | None:
    """
    level=3 -> 'F32.2' -> 'F32'
    level='full' -> keep full 'F32.2'
    """
    if code is None or code is pd.NA or (isinstance(code, float) and pd.isna(code)):
        return None

    c = str(code).strip().upper()
    if c in {"", "NONE", "NAN", "NULL", "<NA>"}:
        return None

    m = re.match(r"^([A-Z])([0-9]{2})(?:\.(\d+))?$", c)
    if not m:
        return None

    letter, d2, dec = m.group(1), m.group(2), m.group(3)
    return f"{letter}{d2}" if level == 3 else c


def map_primary_to_table2(dx_full: str, *, f43_fallback: str = "other", missing_label: str = "Missing") -> str:
    """
    Map ONE full ICD code (e.g., 'F43.1') to a Table-2 category.
    Always returns a STRING (Docx/CatBoost safe; never np.nan).
    """
    if dx_full is None:
        return TABLE2_CATS["missing"]

    s = str(dx_full).strip().upper()
    if s in {"", "<NA>", "NAN", "NONE", "NULL", missing_label.upper()}:
        return TABLE2_CATS["missing"]

    dx3 = norm_icd(s, level=3)
    if dx3 is None:
        return TABLE2_CATS["other"]

    # F10–F19 substances
    if dx3.startswith("F"):
        try:
            fnum = int(dx3[1:3])
        except ValueError:
            fnum = None

        if fnum is not None and 10 <= fnum <= 19:
            return TABLE2_CATS["substance"]
        if fnum is not None and 20 <= fnum <= 29:
            return TABLE2_CATS["psychotic"]
        if fnum is not None and 30 <= fnum <= 39:
            return TABLE2_CATS["affective"]

    # anxiety + OCD
    if dx3 in {"F40", "F41"}:
        return TABLE2_CATS["anxiety"]
    if dx3 == "F42":
        return TABLE2_CATS["ocd"]

    # F43 split by decimals
    if dx3 == "F43":
        if re.match(r"^F43\.0(\D|$)|^F43\.1(\D|$)", s):
            return TABLE2_CATS["asr_ptsd"]
        if re.match(r"^F43\.2(\D|$)", s):
            return TABLE2_CATS["adjustment"]
        return TABLE2_CATS.get(f43_fallback, TABLE2_CATS["other"])

    # somatoform, eating, personality, ADHD
    if dx3 == "F45":
        return TABLE2_CATS["somatoform"]
    if dx3 == "F50":
        return TABLE2_CATS["eating"]
    if dx3 in {"F60", "F61", "F62"}:
        return TABLE2_CATS["personality"]
    if dx3 == "F90":
        return TABLE2_CATS["adhd"]

    return TABLE2_CATS["other"]


def infer_diag_col(df: pd.DataFrame, candidates: list[str]) -> str | None:
    """Return first matching diagnosis column name, else None."""
    for c in candidates:
        if c in df.columns:
            return c
    return None


def map_primary_list_to_table2_list(
    primary_list: list[str],
    *,
    allowed_letter: str = "F",
    missing_label: str = TABLE2_CATS["missing"],
    f43_fallback: str = "other",
) -> list[str]:
    """
    Multi-label mapping: list of ICD codes -> list of Table-2 categories (unique).
    Filters to allowed_letter (default 'F') for psychiatric Table-2 categories.
    """
    if not isinstance(primary_list, list) or len(primary_list) == 0:
        return [missing_label]

    cats: list[str] = []
    for code in primary_list:
        c_full = norm_icd(code, level="full")  # keep decimals for F43 split
        if c_full is None:
            continue
        if allowed_letter is not None and not str(c_full).startswith(allowed_letter):
            continue
        cats.append(map_primary_to_table2(c_full, f43_fallback=f43_fallback, missing_label="Missing"))

    if len(cats) == 0:
        return [missing_label]

    return sorted(set(cats))


def add_primary_table2_multilabel_one_row(df_pat: pd.DataFrame) -> pd.DataFrame:
    """
    Assumes ONE row per patient. Creates/overwrites df_pat['primary_dx_table2']
    as LIST of Table-2 categories per patient (multi-label).
    """
    df_pat = df_pat.copy()
    primary_col = infer_diag_col(df_pat, ["primarydiagnosis_tp", "primarydiagnosis_tp_t0", "primarydiagnosis"])
    if primary_col is None:
        # no diagnosis available
        df_pat["primary_dx_table2"] = [[TABLE2_CATS["missing"]] for _ in range(len(df_pat))]
        return df_pat

    df_pat["_primary_list_raw"] = df_pat[primary_col].apply(parse_icd_codes)
    df_pat["primary_dx_table2"] = df_pat["_primary_list_raw"].apply(
        lambda lst: map_primary_list_to_table2_list(lst, allowed_letter="F", f43_fallback="other")
    )

    # ensure list + fallback to Missing diagnosis
    df_pat["primary_dx_table2"] = df_pat["primary_dx_table2"].apply(
        lambda x: x if isinstance(x, list) and len(x) > 0 else [TABLE2_CATS["missing"]]
    )

    df_pat.drop(columns=["_primary_list_raw"], inplace=True)
    return df_pat


# =============================================================================
# Helpers: formatting + borders
# =============================================================================
def fmt_mean_sd(x: pd.Series) -> str:
    x = pd.to_numeric(x, errors="coerce")
    m = x.mean(skipna=True)
    sd = x.std(skipna=True, ddof=1)
    if pd.isna(m):
        return ""
    return f"{m:.2f} ({sd:.2f})"


def fmt_n_pct(n: int, denom: int) -> str:
    if denom <= 0:
        return f"{n} (0.0%)"
    pct = (n / denom) * 100.0
    return f"{n} ({pct:.1f}%)"


def missing_n_pct(x: pd.Series, N: int) -> str:
    miss = int(pd.isna(x).sum())
    return fmt_n_pct(miss, N)


def infer_id_col(df: pd.DataFrame, candidates: list[str]) -> str:
    for c in candidates:
        if c in df.columns:
            return c
    raise KeyError(
        f"None of the ID columns found: {candidates}. "
        f"Available (first 30): {list(df.columns)[:30]} ..."
    )


def load_df_all() -> pd.DataFrame:
    in_path = Path(DATA_PATH) / "toni_preprocessed.csv"
    if not in_path.exists():
        raise FileNotFoundError(f"Could not find: {in_path}")
    return pd.read_csv(in_path)


def _set_cell_text(
    cell,
    text: str,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    left_indent_in: float | None = None,
):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic

    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if left_indent_in is not None:
        p.paragraph_format.left_indent = Inches(left_indent_in)


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)
        tag = qn(f"w:{edge}")
        element = tcBorders.find(tag)
        if edge_data is None:
            if element is not None:
                tcBorders.remove(element)
            continue
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tcBorders.append(element)
        for key, val in edge_data.items():
            element.set(qn(f"w:{key}"), str(val))


def apply_apa_table_rules(table):
    n_rows = len(table.rows)
    n_cols = len(table.columns)

    for r in range(n_rows):
        for c in range(n_cols):
            _set_cell_border(table.cell(r, c), top=None, bottom=None, left=None, right=None)

    rule = {"val": "single", "sz": "8", "color": "000000"}  # ~1pt

    # header top+bottom
    for c in range(n_cols):
        _set_cell_border(table.cell(0, c), top=rule, bottom=rule)

    # table bottom
    last = n_rows - 1
    for c in range(n_cols):
        _set_cell_border(table.cell(last, c), bottom=rule)


# =============================================================================
# Codebooks (subcategories EN)
# =============================================================================
GENDER = {
    1: "Female",
    2: "Non-binary",
    3: "Genderfluid",
    4: "Male",
    5: "Other",
    6: "Prefer not to say",
}

SCHOOL = {
    0: "No school-leaving certificate",
    1: "Lower secondary school or equivalent",
    2: "Intermediate secondary school or equivalent",
    3: "Higher education entrance qualification or equivalent",
    4: "Completed vocational training",
    5: "Other school-leaving certificate",
}

EMPLOYMENT = {
    1: "Employed",
    2: "Self-employed",
    3: "Unemployed",
    4: "Retired",
    5: "Student / in training",
    6: "Other",
}

LOCATION = {
    1: "Rural",
    2: "Medium-sized city/town",
    3: "Large city",
}

LABEL = {
    "t0_age_p": "Age (years)",
    "t0_gender_p": "Gender identity",
    "t0_distress_sum": "Baseline distress (sum score)",
    "t1_distress_sum": "Week 6 distress (sum score)",
    # multi-label list per patient
    "primary_dx_table2": "Primary diagnosis (ICD-10; Table 2 categories)",
    "t0_school_p": "Highest educational attainment",
    "t0_employment_p": "Employment status",
    "t0_location_p": "Residential location",
    "t0_relationship_p": "Relationship status",
    "t0_migration_p_1": "Migration background",
    "t0_ses_p": "Socioeconomic status",
}

CODEBOOK = {
    "t0_gender_p": GENDER,
    "t0_school_p": SCHOOL,
    "t0_employment_p": EMPLOYMENT,
    "t0_location_p": LOCATION,
}


# =============================================================================
# Build rows for the table
# =============================================================================
def add_continuous_rows(rows, df, col, N):
    label = LABEL.get(col, col)
    rows.append(
        {"Variable": label, "Value": fmt_mean_sd(df[col]), "Missing": missing_n_pct(df[col], N), "is_header": False, "indent": 0.0}
    )


def add_categorical_rows(rows, df, col, N):
    label = LABEL.get(col, col)
    rows.append({"Variable": label, "Value": "", "Missing": missing_n_pct(df[col], N), "is_header": True, "indent": 0.0})

    ser = pd.to_numeric(df[col], errors="coerce")
    denom = int(ser.notna().sum())
    mapper = CODEBOOK.get(col)

    if mapper is not None:
        present_codes = set(ser.dropna().astype(int).unique())
        cat_values = [k for k in mapper.keys() if k in present_codes]
        unexpected = sorted(present_codes - set(mapper.keys()))
        cat_values += unexpected
    else:
        cat_values = sorted(ser.dropna().unique())

    for v in cat_values:
        n = int((ser == v).sum())
        cat_label = mapper.get(int(v), f"Code {int(v)}") if mapper is not None else str(v)
        rows.append({"Variable": cat_label, "Value": fmt_n_pct(n, denom), "Missing": "", "is_header": False, "indent": 0.25})


def detect_binary(series: pd.Series) -> bool:
    s = pd.to_numeric(series, errors="coerce")
    uniq = set(s.dropna().unique().tolist())
    return len(uniq) > 0 and uniq.issubset({0, 1})


def add_binary_heading_yes(rows, df, col, N):
    label = LABEL.get(col, col)
    rows.append({"Variable": label, "Value": "", "Missing": missing_n_pct(df[col], N), "is_header": True, "indent": 0.0})

    ser = pd.to_numeric(df[col], errors="coerce")
    denom = int(ser.notna().sum())
    yes_n = int((ser == 1).sum())
    rows.append({"Variable": "Yes", "Value": fmt_n_pct(yes_n, denom), "Missing": "", "is_header": False, "indent": 0.25})


def add_multilabel_table2_rows(rows, df, col, N, order: list[str]):
    """
    Multi-label categorical: df[col] contains LISTS of Table-2 categories per patient.
    Count patients with category ∈ list. Denominator = N (patients).
    """
    label = LABEL.get(col, col)

    def _is_missing_list(x):
        if not isinstance(x, list) or len(x) == 0:
            return True
        xs = [str(v).strip() for v in x if str(v).strip() != ""]
        if len(xs) == 0:
            return True
        return set(xs).issubset({TABLE2_CATS["missing"]})

    miss_n = int(df[col].apply(_is_missing_list).sum())
    rows.append({"Variable": label, "Value": "", "Missing": fmt_n_pct(miss_n, N), "is_header": True, "indent": 0.0})

    for cat in order:
        n = int(df[col].apply(lambda lst: isinstance(lst, list) and (cat in lst)).sum())
        rows.append({"Variable": cat, "Value": fmt_n_pct(n, N), "Missing": "", "is_header": False, "indent": 0.25})


def make_rows(df: pd.DataFrame) -> list[dict]:
    N = len(df)
    rows: list[dict] = []

    ordered = [
        "t0_age_p",
        "t0_gender_p",
        "t0_distress_sum",
        "t1_distress_sum",
        "primary_dx_table2",  # multi-label list
        "t0_school_p",
        "t0_employment_p",
        "t0_location_p",
        "t0_relationship_p",
        "t0_migration_p_1",
        "t0_ses_p",
    ]
    cols = [c for c in ordered if c in df.columns]

    continuous = {"t0_age_p", "t0_distress_sum", "t1_distress_sum", "t0_ses_p"}
    categorical = {"t0_gender_p", "t0_school_p", "t0_employment_p", "t0_location_p"}
    multilabel_table2 = {"primary_dx_table2"}

    for col in cols:
        if col in continuous:
            add_continuous_rows(rows, df, col, N)
        elif col in categorical:
            add_categorical_rows(rows, df, col, N)
        elif col in multilabel_table2:
            add_multilabel_table2_rows(rows, df, col, N, order=TABLE2_ORDER)
        else:
            if detect_binary(df[col]):
                add_binary_heading_yes(rows, df, col, N)
            else:
                add_continuous_rows(rows, df, col, N)

    return rows


# =============================================================================
# Create DOCX
# =============================================================================
def create_docx_table(rows: list[dict], N: int, out_path: str):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p_num = doc.add_paragraph()
    r1 = p_num.add_run("Table 1")
    r1.italic = True

    p_title = doc.add_paragraph()
    r2 = p_title.add_run("Patient characteristics")
    r2.italic = True

    doc.add_paragraph()

    table = doc.add_table(rows=len(rows) + 1, cols=3)
    table.autofit = True

    _set_cell_text(table.cell(0, 0), "Variable", bold=True)
    _set_cell_text(table.cell(0, 1), "M (SD) / n (%)", bold=True, align="right")
    _set_cell_text(table.cell(0, 2), "Missing n (%)", bold=True, align="right")

    for i, row in enumerate(rows, start=1):
        _set_cell_text(
            table.cell(i, 0),
            row["Variable"],
            bold=row.get("is_header", False),
            align="left",
            left_indent_in=row.get("indent", 0.0),
        )
        _set_cell_text(table.cell(i, 1), row["Value"], align="right")
        _set_cell_text(table.cell(i, 2), row["Missing"], align="right")

    apply_apa_table_rules(table)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        f"Note. N = {N}. Continuous variables are presented as M (SD). "
        f"Categorical variables are presented as n (%). Percentages for binary variables "
        f"refer to the proportion of 'Yes' among non-missing responses. "
        f"Primary diagnosis categories (ICD-10; Table 2) are not mutually exclusive; patients may be counted in multiple categories."
    )

    doc.save(out_path)


def main():
    out_path = "Table_Patient_Features_APA7_updated.docx"

    df_all = load_df_all()
    PAT_ID = infer_id_col(df_all, ["patient_id", "id", "pat_id", "case_id"])

    # If there is indeed one row per patient, this is just a safeguard:
    df_pat = df_all.drop_duplicates(subset=[PAT_ID]).copy()

    # Create multi-label primary diagnosis Table-2 list directly from this row
    df_pat = add_primary_table2_multilabel_one_row(df_pat)

    N = len(df_pat)
    rows = make_rows(df_pat)
    create_docx_table(rows, N, out_path)

    print(f"Wrote: {out_path} (N={N}, source={Path(DATA_PATH) / 'toni_preprocessed.csv'})")


if __name__ == "__main__":
    main()